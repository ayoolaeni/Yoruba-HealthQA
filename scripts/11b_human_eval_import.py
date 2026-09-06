#!/usr/bin/env python
"""Phase 7 -- Import completed human evaluation sheets, unblind, aggregate,
compute agreement (build spec section 9).

Reads every reports/human_eval/rater_*.xlsx (or --rater-files), unblinds
using --blinding-key (never shipped to raters), and emits:
  table_4_7_human_ratings        -- mean Likert score per system per
                                     dimension + harm category counts.
  table_4_8_agreement            -- Fleiss' kappa (harm, categorical) +
                                     Krippendorff's alpha (each ordinal
                                     dimension) + Friedman test across
                                     systems per dimension.
  table_4_9_harm_incidents       -- every individual Moderate/Severe rating,
                                     verbatim (a headline result, not an
                                     appendix, per build spec).
  fig_4_4_human_ratings.png      -- mean Likert score by system, grouped by
                                     dimension.

Usage:
    python scripts/11b_human_eval_import.py --rater-files reports/human_eval/rater_*.xlsx \
        --blinding-key reports/.blinding_key.json
"""
from __future__ import annotations

import argparse
import glob
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.eval.agreement import build_category_count_matrix, fleiss_kappa, krippendorff_alpha_ordinal
from src.eval.human_eval import (
    HARM_CATEGORIES,
    LIKERT_DIMENSIONS,
    harm_distribution_by_system,
    harm_incidents,
    mean_ratings_by_system,
    unblind_rater_rows,
)
from src.utils.logging import MissingInputError, get_logger
from src.utils.manifest import RunManifest
from src.utils.seeding import seed_everything

logger = get_logger("11b_human_eval_import")

SHEET_COLUMNS = (["sample_id", "topic", "question_yo", "slot", "response_yo"]
                  + LIKERT_DIMENSIONS + ["potential_for_harm", "rater_comment"])


def read_rater_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows_iter)]
    rows = []
    for values in rows_iter:
        row = dict(zip(header, values))
        if row.get("sample_id"):
            rows.append({k: ("" if v is None else v) for k, v in row.items()})
    return rows


def validate_rows(rows: list[dict], rater_id: str) -> list[str]:
    problems = []
    for row in rows:
        for dim in LIKERT_DIMENSIONS:
            val = row.get(dim)
            try:
                val_int = int(val)
                if not 1 <= val_int <= 5:
                    raise ValueError
            except (TypeError, ValueError):
                problems.append(f"[{rater_id}] sample={row.get('sample_id')} slot={row.get('slot')}: "
                                 f"'{dim}' must be an integer 1-5, got {val!r}")
        harm = str(row.get("potential_for_harm", "")).strip()
        if harm not in HARM_CATEGORIES:
            problems.append(f"[{rater_id}] sample={row.get('sample_id')} slot={row.get('slot')}: "
                             f"'potential_for_harm' must be one of {HARM_CATEGORIES}, got {harm!r}")
    return problems


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rater-files", default="reports/human_eval/rater_*.xlsx")
    parser.add_argument("--blinding-key", default="reports/.blinding_key.json")
    parser.add_argument("--reports-dir", default="reports")
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()

    blinding_key_path = Path(args.blinding_key)
    if not blinding_key_path.exists():
        raise MissingInputError(str(blinding_key_path), stage="11b_human_eval_import",
                                 hint="run scripts/11_human_eval_export.py first")
    with open(blinding_key_path, encoding="utf-8") as f:
        blinding_map = json.load(f)

    rater_paths = sorted(Path(p) for p in glob.glob(args.rater_files))
    if not rater_paths:
        raise MissingInputError(args.rater_files, stage="11b_human_eval_import",
                                 hint="no completed rater workbooks found")

    seed = seed_everything(args.seed)

    all_problems = []
    all_ratings = []
    for path in rater_paths:
        rater_id = path.stem
        rows = read_rater_xlsx(path)
        problems = validate_rows(rows, rater_id)
        all_problems.extend(problems)
        if not problems:
            all_ratings.extend(unblind_rater_rows(rows, rater_id=rater_id, blinding_map=blinding_map))

    if all_problems:
        for p in all_problems[:50]:
            logger.error(p)
        raise ValueError(f"{len(all_problems)} problem(s) found across rater sheets -- fix and re-run")

    logger.info(f"loaded {len(all_ratings)} individual ratings from {len(rater_paths)} raters")

    import numpy as np
    import pandas as pd

    with RunManifest(stage="11b_human_eval_import", config={"n_raters": len(rater_paths), "n_ratings": len(all_ratings)},
                      seed=seed) as run:
        reports_dir = Path(args.reports_dir)
        reports_dir.mkdir(parents=True, exist_ok=True)

        # --- table_4_7: mean ratings + harm distribution ---
        means = mean_ratings_by_system(all_ratings)
        harm_dist = harm_distribution_by_system(all_ratings)
        rows_47 = []
        for system_id in means:
            row = {"system": system_id, **{dim: round(v, 3) for dim, v in means[system_id].items()}}
            for cat in HARM_CATEGORIES:
                row[f"harm_{cat.lower()}"] = harm_dist.get(system_id, {}).get(cat, 0)
            rows_47.append(row)
        df_47 = pd.DataFrame(rows_47)
        df_47.to_csv(reports_dir / "table_4_7_human_ratings.csv", index=False)
        with open(reports_dir / "table_4_7_human_ratings.md", "w", encoding="utf-8") as f:
            f.write(df_47.to_markdown(index=False))
        run.record_output(str(reports_dir / "table_4_7_human_ratings.csv"))

        # --- table_4_8: agreement ---
        by_item_harm: dict[tuple, list[str]] = {}
        by_item_ordinal: dict[str, dict[tuple, dict[str, float]]] = {dim: {} for dim in LIKERT_DIMENSIONS}
        raters = sorted({r.rater_id for r in all_ratings})
        for r in all_ratings:
            key = (r.sample_id, r.system_id)
            by_item_harm.setdefault(key, []).append(r.harm)
            for dim in LIKERT_DIMENSIONS:
                by_item_ordinal[dim].setdefault(key, {})[r.rater_id] = r.ratings[dim]

        agreement_rows = []
        # Fleiss' kappa needs a uniform rater count per item -- keep only
        # items every rater actually rated.
        complete_items = [key for key, harms in by_item_harm.items() if len(harms) == len(raters)]
        if len(complete_items) >= 2 and len(raters) >= 2:
            matrix = build_category_count_matrix([by_item_harm[k] for k in complete_items], HARM_CATEGORIES)
            kappa = fleiss_kappa(matrix)
            agreement_rows.append({"dimension": "potential_for_harm", "statistic": "fleiss_kappa", "value": round(kappa, 4)})
        else:
            logger.warning("not enough fully double-rated items/raters for Fleiss' kappa -- skipped")

        for dim in LIKERT_DIMENSIONS:
            item_keys = list(by_item_ordinal[dim].keys())
            reliability_data = [
                [by_item_ordinal[dim][key].get(rater_id) for key in item_keys]
                for rater_id in raters
            ]
            if len(raters) >= 2 and len(item_keys) >= 2:
                alpha = krippendorff_alpha_ordinal(reliability_data)
                agreement_rows.append({"dimension": dim, "statistic": "krippendorff_alpha", "value": round(alpha, 4)})

        # Friedman test: for each dimension, is there a significant
        # difference across systems? One block per sample_id (averaged
        # across raters for that sample+system), matched across systems.
        systems = sorted({r.system_id for r in all_ratings})
        for dim in LIKERT_DIMENSIONS:
            per_system_per_sample: dict[str, dict[str, list[float]]] = {s: {} for s in systems}
            for r in all_ratings:
                per_system_per_sample[r.system_id].setdefault(r.sample_id, []).append(r.ratings[dim])
            common_samples = set.intersection(*[set(per_system_per_sample[s]) for s in systems]) if systems else set()
            if len(common_samples) < 3 or len(systems) < 3:
                continue  # Friedman requires >=3 related samples and is only meaningful for >=3 systems
            common_samples = sorted(common_samples)
            blocks = [[float(np.mean(per_system_per_sample[s][sample_id])) for s in systems]
                      for sample_id in common_samples]
            from scipy.stats import friedmanchisquare

            columns = list(zip(*blocks))
            stat, p_value = friedmanchisquare(*columns)
            agreement_rows.append({"dimension": dim, "statistic": "friedman_chi2", "value": round(float(stat), 4)})
            agreement_rows.append({"dimension": dim, "statistic": "friedman_p_value", "value": round(float(p_value), 6)})

        df_48 = pd.DataFrame(agreement_rows)
        df_48.to_csv(reports_dir / "table_4_8_agreement.csv", index=False)
        with open(reports_dir / "table_4_8_agreement.md", "w", encoding="utf-8") as f:
            f.write(df_48.to_markdown(index=False) if len(df_48) else "(no agreement statistics computed)")
        run.record_output(str(reports_dir / "table_4_8_agreement.csv"))

        # --- table_4_9: harm incidents ---
        incidents = harm_incidents(all_ratings)
        df_49 = pd.DataFrame(incidents)
        df_49.to_csv(reports_dir / "table_4_9_harm_incidents.csv", index=False)
        with open(reports_dir / "table_4_9_harm_incidents.md", "w", encoding="utf-8") as f:
            f.write(df_49.to_markdown(index=False) if len(df_49) else "(no Moderate/Severe incidents reported)")
        run.record_output(str(reports_dir / "table_4_9_harm_incidents.csv"))
        logger.info(f"{len(incidents)} individual Moderate/Severe harm incidents recorded")

        # --- fig_4_4: mean Likert ratings by system ---
        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt

            fig, ax = plt.subplots(figsize=(9, 5), dpi=300)
            width = 0.8 / max(len(means), 1)
            x_base = range(len(LIKERT_DIMENSIONS))
            for i, system_id in enumerate(sorted(means)):
                values = [means[system_id].get(dim, 0) for dim in LIKERT_DIMENSIONS]
                ax.bar([x + i * width for x in x_base], values, width, label=system_id)
            ax.set_xticks([x + width * (len(means) - 1) / 2 for x in x_base])
            ax.set_xticklabels(LIKERT_DIMENSIONS, rotation=20, ha="right")
            ax.set_ylabel("mean Likert score (1-5)")
            ax.set_title("Human evaluation: mean ratings by system")
            ax.legend()
            fig.tight_layout()
            fig_path = reports_dir / "fig_4_4_human_ratings.png"
            fig.savefig(fig_path, dpi=300)
            plt.close(fig)
            run.record_output(str(fig_path))
        except ImportError:
            logger.warning("matplotlib not installed -- skipped fig_4_4_human_ratings.png")

        logger.info("human evaluation import/aggregation complete")


if __name__ == "__main__":
    main()

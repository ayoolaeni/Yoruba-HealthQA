#!/usr/bin/env python
"""Phase 9 -- Error analysis on system M's lowest-scoring responses
(build spec section 11).

Classifying WHY a response is wrong is a human judgment call (see
src/eval/error_taxonomy.py docstring) -- this script does not auto-label
factual_error/terminological_error/hallucination/refusal_failure. It:

  --export: selects the --n lowest per-sentence-chrF++ scoring M responses
      from reports/generations/M.jsonl, pre-fills SUGGESTED diacritic_error /
      code_switching / omission / disfluency flags from cheap heuristics, and
      writes an error-analysis spreadsheet for a human reviewer to correct
      the suggestions and fill in the four judgment-only categories.
  --aggregate: reads the completed spreadsheet and emits
      table_4_10_error_taxonomy (counts by error type x topic) and
      fig_4_6_errors_by_topic.png.

Usage:
    python scripts/12_error_analysis.py --export --test-jsonl data/final/test.jsonl \
        --generations reports/generations/M.jsonl --n 100
    python scripts/12_error_analysis.py --aggregate --reviewed reports/error_analysis_reviewed.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.data.schema import read_jsonl
from src.eval.error_taxonomy import ERROR_CATEGORIES, select_lowest_scoring, suggest_flags
from src.eval.metrics import sentence_chrf_scores
from src.utils.logging import MissingInputError, get_logger
from src.utils.manifest import RunManifest
from src.utils.seeding import seed_everything

logger = get_logger("12_error_analysis")

SHEET_COLUMNS = (["id", "topic", "question_yo", "reference", "hypothesis", "chrf++"]
                  + ERROR_CATEGORIES + ["reviewer_notes"])


def build_export_rows(test_records: dict[str, dict], generations: list[dict], n: int) -> list[dict]:
    scored = []
    hyps = [g["hypothesis"] for g in generations]
    refs = [test_records[g["id"]]["output"] for g in generations]
    chrf_scores = sentence_chrf_scores(hyps, refs)
    for g, score in zip(generations, chrf_scores):
        scored.append({**g, "chrf++": score, "topic": test_records[g["id"]].get("topic", "unclassified")})

    worst = select_lowest_scoring(scored, "chrf++", n)
    rows = []
    for item in worst:
        flags = suggest_flags(item["hypothesis"], item["reference"] if "reference" in item else test_records[item["id"]]["output"])
        row = {
            "id": item["id"], "topic": item["topic"], "question_yo": item["instruction"],
            "reference": test_records[item["id"]]["output"], "hypothesis": item["hypothesis"],
            "chrf++": round(item["chrf++"], 2),
            "factual_error": "", "terminological_error": "", "refusal_failure": "", "hallucination": "",
            "diacritic_error": "SUGGESTED" if flags.diacritic_error else "",
            "code_switching": "SUGGESTED" if flags.code_switching else "",
            "omission": "SUGGESTED" if flags.omission else "",
            "disfluency": "SUGGESTED" if flags.disfluency else "",
            "reviewer_notes": "",
        }
        rows.append(row)
    return rows


def write_xlsx(rows: list[dict], out_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "error_analysis"
    ws.append(SHEET_COLUMNS)
    for row in rows:
        ws.append([row[c] for c in SHEET_COLUMNS])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def read_xlsx_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows_iter)]
    rows = []
    for values in rows_iter:
        row = dict(zip(header, values))
        if row.get("id"):
            rows.append({k: ("" if v is None else v) for k, v in row.items()})
    return rows


def cmd_export(args) -> None:
    test_path = Path(args.test_jsonl)
    gen_path = Path(args.generations)
    for p in (test_path, gen_path):
        if not p.exists():
            raise MissingInputError(str(p), stage="12_error_analysis")

    seed = seed_everything(args.seed)
    test_records = {r["id"]: r for r in read_jsonl(test_path)}
    generations = list(read_jsonl(gen_path))
    missing = [g["id"] for g in generations if g["id"] not in test_records]
    if missing:
        raise ValueError(f"{len(missing)} generation ids not found in {test_path}")

    with RunManifest(stage="12_error_analysis_export", config={"n": args.n}, seed=seed) as run:
        rows = build_export_rows(test_records, generations, args.n)
        out_path = Path(args.out)
        write_xlsx(rows, out_path)
        run.record_output(str(out_path))
        logger.info(f"exported {len(rows)} lowest-scoring responses for review -> {out_path}")


def cmd_aggregate(args) -> None:
    reviewed_path = Path(args.reviewed)
    if not reviewed_path.exists():
        raise MissingInputError(str(reviewed_path), stage="12_error_analysis", hint="run --export first, "
                                 "have a reviewer confirm/correct the suggested flags and fill in judgment columns")

    seed = seed_everything(args.seed)
    rows = read_xlsx_rows(reviewed_path)

    # A cell still reading the export's "SUGGESTED" sentinel means the human
    # reviewer never actually looked at it -- counting it as a confirmed
    # error would let an unreviewed heuristic stand in for human judgment,
    # exactly what this tool exists to prevent (build-spec rule 1). Refuse to
    # aggregate until every suggestion has been explicitly resolved to a
    # confirmed value or cleared.
    unresolved = [(row.get("id"), cat) for row in rows for cat in ERROR_CATEGORIES
                  if str(row.get(cat, "")).strip().upper() == "SUGGESTED"]
    if unresolved:
        for rid, cat in unresolved[:50]:
            logger.error(f"unresolved suggested flag: id={rid} category={cat}")
        raise ValueError(f"{len(unresolved)} SUGGESTED flag(s) in {reviewed_path} were never confirmed or "
                          f"cleared by a reviewer -- resolve every SUGGESTED cell (to TRUE/YES/1/X to confirm, "
                          f"or clear it to reject) before aggregating")

    import pandas as pd

    with RunManifest(stage="12_error_analysis_aggregate", config={"reviewed": str(reviewed_path)}, seed=seed) as run:
        counts = []
        for category in ERROR_CATEGORIES:
            by_topic: dict[str, int] = {}
            for row in rows:
                value = str(row.get(category, "")).strip().upper()
                if value in ("TRUE", "YES", "1", "X"):
                    topic = row.get("topic", "unclassified")
                    by_topic[topic] = by_topic.get(topic, 0) + 1
            for topic, n in by_topic.items():
                counts.append({"error_type": category, "topic": topic, "count": n})

        df = pd.DataFrame(counts) if counts else pd.DataFrame(columns=["error_type", "topic", "count"])
        reports_dir = Path(args.reports_dir)
        reports_dir.mkdir(parents=True, exist_ok=True)
        csv_p = reports_dir / "table_4_10_error_taxonomy.csv"
        md_p = reports_dir / "table_4_10_error_taxonomy.md"
        df.to_csv(csv_p, index=False)
        with open(md_p, "w", encoding="utf-8") as f:
            f.write(df.to_markdown(index=False) if len(df) else "(no confirmed error flags in reviewed sheet)")
        run.record_output(str(csv_p))

        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt

            if len(df):
                pivot = df.pivot_table(index="topic", columns="error_type", values="count", fill_value=0)
                fig, ax = plt.subplots(figsize=(9, 5), dpi=300)
                pivot.plot(kind="bar", stacked=True, ax=ax)
                ax.set_ylabel("count")
                ax.set_title("Error types by topic (system M, lowest-scoring responses)")
                fig.tight_layout()
                fig_path = reports_dir / "fig_4_6_errors_by_topic.png"
                fig.savefig(fig_path, dpi=300)
                plt.close(fig)
                run.record_output(str(fig_path))
            else:
                logger.warning("no confirmed error flags -- skipped fig_4_6_errors_by_topic.png")
        except ImportError:
            logger.warning("matplotlib not installed -- skipped fig_4_6_errors_by_topic.png")

        logger.info(f"wrote error taxonomy table ({len(counts)} (error_type, topic) rows) -> {csv_p}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--test-jsonl", default="data/final/test.jsonl")
    parser.add_argument("--generations", default="reports/generations/M.jsonl")
    parser.add_argument("--n", type=int, default=100)
    parser.add_argument("--out", default="reports/error_analysis_for_review.xlsx")
    parser.add_argument("--reviewed", default="reports/error_analysis_reviewed.xlsx")
    parser.add_argument("--reports-dir", default="reports")
    parser.add_argument("--seed", type=int, default=42)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--export", action="store_true")
    mode.add_argument("--aggregate", action="store_true")
    args = parser.parse_args()

    if args.export:
        cmd_export(args)
    else:
        cmd_aggregate(args)


if __name__ == "__main__":
    main()
